"""XLSX emitter: the data-truth artifact.

Role split inside the standard:
  HTML  the reading artifact  - full fidelity, opens anywhere, no license
  XLSX  the data artifact     - every number printed in the HTML must be here,
                                as a real cell in a real range you can pivot

So this emitter is deliberately not a screenshot of the HTML. Prose collapses
into one summary sheet; tables, matrices, and heat grids become real ranges with
frozen headers, autofilters, and number formats.

On determinism, honestly: an .xlsx is a zip, and zip entries carry timestamps,
so the file bytes are not stable across runs. Byte-comparing it would be a test
that fails for the wrong reason. `content_digest()` hashes what actually matters
-- sheet names, cell coordinates, and values -- and that is what the golden test
compares. The HTML emitter is the one held to byte equality.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .. import ir as IR

# Palette mirrors sds-console so the two artifacts read as one system.
INK = "FF0F1B22"
INK2 = "FF465A62"
ACCENT = "FF1C5CAB"
RULE = "FFD7DFE1"
SUNK = "FFE8EDEF"
GOOD, WARN, CRIT, INFO, NA = "FF0E7A55", "FFA86F00", "FFC23A39", "FF1C5CAB", "FF6B7F88"
STATE_FILL = {
    "ok": "FFE3F4ED", "warn": "FFFBF2DC", "crit": "FFFBE9E8",
    "info": "FFE6EEF8", "na": "FFEEF2F4",
}
STATE_FONT = {"ok": GOOD, "warn": WARN, "crit": CRIT, "info": INFO, "na": NA}

FONT_NAME = "IBM Plex Sans KR"
MAX_SHEET_NAME = 31
MAX_COL_WIDTH = 52


def _sheet_name(used: set[str], raw: str) -> str:
    """Excel sheet names: <=31 chars, no []:*?/\\ , unique."""
    s = str(raw)
    for ch in "[]:*?/\\":
        s = s.replace(ch, " ")
    s = " ".join(s.split())[:MAX_SHEET_NAME] or "Sheet"
    base, i = s, 2
    while s in used:
        suffix = f" ({i})"
        s = base[: MAX_SHEET_NAME - len(suffix)] + suffix
        i += 1
    used.add(s)
    return s


def _cell_value(v: Any) -> Any:
    """Unwrap badge dicts; keep numbers numeric so Excel can aggregate them."""
    if isinstance(v, dict):
        return v.get("label", v.get("state", ""))
    if isinstance(v, (int, float)) or v is None:
        return v
    return str(v)


def _header(ws, row: int, labels: list[str], aligns: list[str]) -> None:
    thin = Side(style="thin", color=RULE)
    accent = Side(style="medium", color=ACCENT)
    for ci, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=ci, value=label)
        c.font = Font(name=FONT_NAME, bold=True, size=10, color=INK2)
        c.fill = PatternFill("solid", fgColor="FFFFFFFF")
        c.border = Border(bottom=accent, left=thin, right=thin)
        c.alignment = Alignment(
            horizontal="right" if aligns[ci - 1] == "num" else "left",
            vertical="center", wrap_text=True,
        )
    ws.row_dimensions[row].height = 26


def _autosize(ws, labels: list[str], rows: list[list], aligns: list[str]) -> None:
    for ci, label in enumerate(labels, start=1):
        width = len(str(label)) + 4
        for r in rows[:400]:                      # sampling 400 rows is enough
            if ci - 1 < len(r):
                width = max(width, len(str(_cell_value(r[ci - 1]))) + 3)
        if aligns[ci - 1] == "wrap":
            width = min(width, 44)
        ws.column_dimensions[get_column_letter(ci)].width = min(width, MAX_COL_WIDTH)


def _write_table(ws, block: dict, start: int) -> int:
    cols = block.get("columns") or []
    rows = block.get("rows") or []
    labels = [str(c.get("label") or c.get("key") or "") for c in cols]
    aligns = [c.get("align", "text") for c in cols]

    r = start
    if block.get("title"):
        c = ws.cell(row=r, column=1, value=str(block["title"]))
        c.font = Font(name=FONT_NAME, bold=True, size=13, color=INK)
        r += 1
    # The grain caption is not decoration; it travels with the data.
    if block.get("row_grain"):
        c = ws.cell(row=r, column=1, value=f"행 단위: {block['row_grain']}")
        c.font = Font(name=FONT_NAME, size=9, italic=True, color=INK2)
        r += 1
    if block.get("caption"):
        c = ws.cell(row=r, column=1, value=str(block["caption"]))
        c.font = Font(name=FONT_NAME, size=9, color=INK2)
        r += 1
    r += 1

    head_row = r
    _header(ws, head_row, labels, aligns)
    thin = Side(style="thin", color=RULE)

    for ri, row in enumerate(rows):
        for ci, col in enumerate(cols, start=1):
            raw = row[ci - 1] if ci - 1 < len(row) else None
            cell = ws.cell(row=head_row + 1 + ri, column=ci, value=_cell_value(raw))
            cell.font = Font(name=FONT_NAME, size=10, color=INK)
            cell.border = Border(bottom=thin)
            al = aligns[ci - 1]
            if al == "num":
                cell.alignment = Alignment(horizontal="right")
                digits = col.get("digits")
                if isinstance(raw, float):
                    cell.number_format = "#,##0." + "0" * (digits if digits is not None else 2)
                elif isinstance(raw, int):
                    cell.number_format = "#,##0"
            elif al == "wrap":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            elif al == "badge":
                state = raw.get("state") if isinstance(raw, dict) else raw
                if state in STATE_FILL:
                    cell.fill = PatternFill("solid", fgColor=STATE_FILL[state])
                    cell.font = Font(name=FONT_NAME, size=10, bold=True,
                                     color=STATE_FONT[state])
                cell.alignment = Alignment(horizontal="center")

    end = head_row + len(rows)
    if rows and cols:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(cols))}{end}"
        ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    _autosize(ws, labels, rows, aligns)
    return end + 2


def _write_grid(ws, block: dict, start: int) -> int:
    """matrix and heatgrid share a shape: row headers + a rectangle."""
    cols = block.get("columns") or []
    rows = block.get("rows") or []
    cells = block.get("cells") or []
    labels = [str(block.get("corner") or "")] + [str(c) for c in cols]
    aligns = ["text"] + ["num" if block.get("type") == "heatgrid" else "wrap"] * len(cols)

    r = start
    if block.get("title"):
        c = ws.cell(row=r, column=1, value=str(block["title"]))
        c.font = Font(name=FONT_NAME, bold=True, size=13, color=INK)
        r += 2

    head_row = r
    _header(ws, head_row, labels, aligns)
    thin = Side(style="thin", color=RULE)
    for ri, rname in enumerate(rows):
        rr = head_row + 1 + ri
        hc = ws.cell(row=rr, column=1, value=str(rname))
        hc.font = Font(name=FONT_NAME, size=10, bold=True, color=INK2)
        hc.fill = PatternFill("solid", fgColor=SUNK)
        hc.border = Border(bottom=thin)
        line = cells[ri] if ri < len(cells) else []
        for ci, v in enumerate(line, start=2):
            cell = ws.cell(row=rr, column=ci, value=_cell_value(v))
            cell.font = Font(name=FONT_NAME, size=10, color=INK)
            cell.border = Border(bottom=thin)
            if isinstance(v, (int, float)):
                cell.alignment = Alignment(horizontal="center")
                cell.number_format = "#,##0"
            else:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = ws.cell(row=head_row + 1, column=2)
    _autosize(ws, labels, [[r] for r in rows], aligns)
    return head_row + len(rows) + 2


def _write_summary(ws, ir: dict) -> None:
    m = ir.get("docmeta") or {}
    r = 1
    c = ws.cell(row=r, column=1, value=str(m.get("title") or ""))
    c.font = Font(name=FONT_NAME, bold=True, size=16, color=INK)
    r += 2
    for key in ("version", "as_of", "org", "author", "approver", "next_review", "classification"):
        if m.get(key):
            k = ws.cell(row=r, column=1, value=key)
            k.font = Font(name=FONT_NAME, size=10, bold=True, color=INK2)
            ws.cell(row=r, column=2, value=str(m[key])).font = Font(name=FONT_NAME, size=10)
            r += 1
    r += 1

    ctx_metrics = ir.get("metrics") or {}

    def flat(text: str) -> str:
        return IR.resolve_metrics(str(text), ctx_metrics)

    for si, section in enumerate(ir.get("sections") or [], start=1):
        h = ws.cell(row=r, column=1, value=f"{si}. {section.get('title') or ''}")
        h.font = Font(name=FONT_NAME, bold=True, size=12, color=ACCENT)
        r += 1
        if section.get("intro"):
            ws.cell(row=r, column=2, value=flat(section["intro"])).font = Font(name=FONT_NAME, size=10)
            r += 1
        for b in section.get("blocks") or []:
            bt = b.get("type")
            if bt == "decision_request":
                for i, it in enumerate(b.get("items") or [], start=1):
                    ws.cell(row=r, column=2, value=f"[요청 {i}] {flat(it.get('ask'))}").font = \
                        Font(name=FONT_NAME, size=10, bold=True, color=CRIT)
                    r += 1
                    ws.cell(row=r, column=3, value=flat(it.get("rationale"))).font = \
                        Font(name=FONT_NAME, size=10, color=INK2)
                    r += 1
            elif bt == "bullets":
                for it in b.get("items") or []:
                    ws.cell(row=r, column=2, value="• " + flat(it.get("text"))).font = \
                        Font(name=FONT_NAME, size=10)
                    r += 1
            elif bt == "narrative":
                for para in b.get("paragraphs") or []:
                    ws.cell(row=r, column=2, value=flat(para)).font = Font(name=FONT_NAME, size=10)
                    r += 1
            elif bt == "chart":
                ws.cell(row=r, column=2, value="[결론] " + flat(b.get("takeaway"))).font = \
                    Font(name=FONT_NAME, size=10, bold=True)
                r += 1
            elif bt == "callout":
                ws.cell(row=r, column=2, value=flat(b.get("body"))).font = \
                    Font(name=FONT_NAME, size=10, italic=True, color=INK2)
                r += 1
            elif bt in ("table", "matrix", "heatgrid", "risks", "changelog", "appendix_source"):
                label = b.get("title") or bt
                ws.cell(row=r, column=2, value=f"→ 시트 참조: {label}").font = \
                    Font(name=FONT_NAME, size=9, color=NA)
                r += 1
        r += 1

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 96
    ws.column_dimensions["C"].width = 96
    ws.sheet_view.showGridLines = False


def _write_metrics(ws, ir: dict) -> None:
    """One row per metric, with its formula. This sheet is the KPI register."""
    labels = ["key", "value", "unit", "formula", "source_sql", "source_rows", "computed_at"]
    _header(ws, 1, labels, ["text", "num", "text", "wrap", "wrap", "num", "text"])
    thin = Side(style="thin", color=RULE)
    for ri, (key, spec) in enumerate(sorted((ir.get("metrics") or {}).items()), start=2):
        vals = [key, spec.get("value"), spec.get("unit"), spec.get("formula"),
                spec.get("source_sql"), spec.get("source_rows"), spec.get("computed_at")]
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=ri, column=ci, value=_cell_value(v))
            c.font = Font(name=FONT_NAME, size=10, color=INK)
            c.border = Border(bottom=thin)
            if ci in (2, 6) and isinstance(v, (int, float)):
                c.alignment = Alignment(horizontal="right")
                c.number_format = "#,##0.##" if isinstance(v, float) else "#,##0"
            if ci in (4, 5):
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in (("A", 26), ("B", 14), ("C", 10), ("D", 52), ("E", 52), ("F", 12), ("G", 22)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def emit(ir: dict, path: str | Path) -> Path:
    IR.assert_valid(ir)
    wb = Workbook()
    used: set[str] = set()

    ws = wb.active
    ws.title = _sheet_name(used, "00_Summary")
    _write_summary(ws, ir)

    if ir.get("metrics"):
        _write_metrics(wb.create_sheet(_sheet_name(used, "01_Metrics")), ir)

    for si, section in enumerate(ir.get("sections") or [], start=1):
        data_blocks = [
            b for b in (section.get("blocks") or [])
            if b.get("type") in ("table", "matrix", "heatgrid", "risks", "changelog", "appendix_source")
        ]
        if not data_blocks:
            continue
        sheet = wb.create_sheet(_sheet_name(used, f"{si:02d}_{section.get('title') or 'Section'}"))
        row = 1
        for b in data_blocks:
            bt = b.get("type")
            if bt in ("matrix", "heatgrid"):
                row = _write_grid(sheet, b, row)
            elif bt == "risks":
                row = _write_table(sheet, _risks_as_table(b), row)
            elif bt == "changelog":
                row = _write_table(sheet, _changelog_as_table(b), row)
            elif bt == "appendix_source":
                row = _write_table(sheet, _sources_as_table(b), row)
            else:
                row = _write_table(sheet, b, row)
        sheet.sheet_view.showGridLines = False

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# --- shape adapters: reuse one table writer for the list-shaped blocks -------
def _risks_as_table(b: dict) -> dict:
    return {
        "type": "table", "title": b.get("title") or "리스크 · 전제",
        "row_grain": "행 1개 = 리스크 1건",
        "columns": [
            {"label": "리스크", "align": "wrap"}, {"label": "영향", "align": "wrap"},
            {"label": "완화", "align": "wrap"}, {"label": "오너", "align": "text"},
            {"label": "상태", "align": "badge"},
        ],
        "rows": [[it.get("risk"), it.get("impact"), it.get("mitigation"), it.get("owner"),
                  {"state": it.get("state") or "info", "label": it.get("state_label") or it.get("state") or "info"}]
                 for it in b.get("items") or []],
    }


def _changelog_as_table(b: dict) -> dict:
    return {
        "type": "table", "title": b.get("title") or "변경 이력",
        "row_grain": "행 1개 = 변경 1건",
        "columns": [{"label": "시점"}, {"label": "변경", "align": "wrap"},
                    {"label": "사유", "align": "wrap"}, {"label": "승인"}],
        "rows": [[it.get("when"), it.get("what"), it.get("why"), it.get("who")]
                 for it in b.get("items") or []],
    }


def _sources_as_table(b: dict) -> dict:
    return {
        "type": "table", "title": b.get("title") or "원천 데이터",
        "row_grain": "행 1개 = 원천 파일 1개",
        "columns": [{"label": "원천"}, {"label": "경로", "align": "wrap"},
                    {"label": "시트"}, {"label": "행수", "align": "num"}, {"label": "기준일"}],
        "rows": [[it.get("name"), it.get("path"), it.get("sheet"), it.get("rows"), it.get("as_of")]
                 for it in b.get("items") or []],
    }


# ---------------------------------------------------------------------------
def content_digest(path: str | Path) -> str:
    """Stable hash of what the workbook says, ignoring zip metadata.

    Use this instead of a byte comparison: .xlsx zip entries carry timestamps,
    so file bytes differ between runs for reasons that have nothing to do with
    the document.
    """
    wb = load_workbook(path, data_only=False)
    payload = []
    for name in wb.sheetnames:
        ws = wb[name]
        cells = []
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    cells.append([c.coordinate, c.value if not isinstance(c.value, float)
                                  else round(c.value, 6)])
        payload.append([name, cells])
    blob = json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha256(blob.encode("utf-8")).hexdigest()
